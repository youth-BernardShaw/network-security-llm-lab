# -*- coding: utf-8 -*-
"""Generate optional4/state_matrix.xlsx (role x object x action x state matrix).

Only depends on openpyxl; run from repo root:
    python optional4/tools/make_state_matrix.py
Cell content is maintained in this script so the workbook is reproducible.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "state_matrix.xlsx"

HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE6F0")
VIOLATED_FILL = PatternFill("solid", fgColor="FBD5D5")
ENFORCED_FILL = PatternFill("solid", fgColor="D8EAD3")
WRAP = Alignment(vertical="top", wrap_text=True)

# (靶场, 模块/代码位置, 角色, 对象, 动作, 前置状态, 后置状态, 服务端强制, 判定依据, 关联用例)
MATRIX = [
    ("sqli-labs", "Less-24/login_create.php:25-48", "R1 匿名(注册)", "O1 security.users(写入)",
     "注册写入(转义后INSERT)", "未注册", "已注册(字面值落库)", "成立-写入",
     "mysql_escape_string 仅保证SQL存储,字符集无约束(含<>'#)", "OPT4-SO-01 第一阶段"),
    ("sqli-labs", "Less-24/login.php:17-40", "R1 匿名", "O1->O1S 会话赋值",
     "登录(SELECT精确匹配)", "已注册", "已登录(会话保存落库原文)", "成立",
     "主体SQLI-005已验证:落库原文原样回到会话", "主体 wp4-20261019-sqli005-low"),
    ("sqli-labs", "Less-24/logged-in.php:26", "R3 已登录", "O1S 会话用户名",
     "渲染会话用户名(echo)", "已登录", "页面HTML输出", "违反",
     "echo $_SESSION[username] 无输出编码,HTML标记载荷原样进入DOM", "OPT4-SO-01 汇点(新增)"),
    ("sqli-labs", "Less-24/pass_change.php:31,38", "R3 已登录", "O1 security.users(更新)",
     "改密(UPDATE拼接会话名)", "已登录", "口令已变", "违反",
     "会话名未经转义拼入UPDATE,主体SQLI-005已验证", "主体 wp4-20261019-sqli005-low"),
    ("sqli-labs", "Less-24/pass_change.php:7-14", "R1 匿名", "O1",
     "无会话直接POST改密", "未登录", "被拒绝(302)", "成立",
     "服务端重定向index.php,状态前置强制", "SEQ-C01 反例"),
    ("DVWA", "xss_s/source/low.php:11-20", "R3 已登录", "O2 guestbook(写入)",
     "留言写入(low)", "输入", "落库字面值", "违反-读取端",
     "仅DB转义,无输出编码路径(主体XSS-005已验)", "主体 wp4-20261019-xss005-low"),
    ("DVWA", "xss_s/source/medium.php:14-16", "R3 已登录", "O2 comment字段",
     "message写入(medium)", "输入", "落库空串(元素被剥除)", "成立",
     "strip_tags剥除整个元素,可执行构造从未持久化", "OPT4-SO-02 同页对照字段"),
    ("DVWA", "xss_s/source/medium.php:21-24", "R3 已登录", "O2 name字段",
     "name写入(medium)", "输入", "落库字面值", "违反",
     "仅str_replace删除小写<script>,事件类标签原样落库", "OPT4-SO-01同类/OPT4-SO-02"),
    ("DVWA", "xss_s/source/impossible.php:3-33", "R3 已登录", "O2",
     "写入(impossible)", "输入", "落库预编码值", "成立",
     "Token+双字段htmlspecialchars+PDO参数化", "反例基线"),
    ("DVWA", "xss_s/index.php(渲染)", "R3 已登录", "O2 guestbook(读取)",
     "读取渲染", "落库", "HTML输出", "违反",
     "输出端对name无上下文编码(low/medium),主体low已验", "OPT4-SO-02 汇点"),
    ("DVWA", "captcha/source/low.php:5-50", "R3 已登录", "O6 step=1",
     "提交CAPTCHA+新口令", "初始", "确认页", "条件成立",
     "reCAPTCHA密钥未配置时本地恒失败,流程依赖外服务", "OPT4-LG-01 合法路径"),
    ("DVWA", "captcha/source/low.php:52-86", "R3 已登录", "O6+O3 口令",
     "直接POST step=2", "任意(跳过step=1)", "口令已变", "违反",
     "step=2分支不校验step=1发生过,无当前口令校验", "OPT4-LG-01a 状态跳跃"),
    ("DVWA", "captcha/source/medium.php:64-69", "R3 已登录", "O6+O3",
     "POST step=2+passed_captcha=true", "任意", "口令已变", "违反",
     "校验对象是客户端隐藏域参数,可伪造", "OPT4-LG-01b 客户端状态标记"),
    ("DVWA", "captcha/source/impossible.php", "R3 已登录", "O6+O3",
     "改密(impossible)", "已通过Token", "口令已变", "成立",
     "Anti-CSRF Token+当前口令+PDO", "反例基线"),
    ("DVWA", "csrf/source/low.php:3-24", "R3 已登录(站内)", "O3 口令",
     "GET方式改密", "已登录", "口令已变", "单步合法",
     "站内用户主动发起时行为合法(无Token/当前口令仍是缺陷)", "OPT4-LG-02 第一步"),
    ("DVWA", "csrf/source/low.php:3-24", "R5受害浏览器+R6攻击页", "O3",
     "跨站GET改密(受害者会话)", "已登录", "口令已变", "违反",
     "无Token/无Origin/无当前口令,GET可被<img>等元素自动触发", "OPT4-LG-02 组合越权"),
    ("DVWA", "csrf/source/impossible.php:3-53", "R5+R6", "O3+O4",
     "跨站改密尝试(impossible)", "已登录", "被拒绝", "成立",
     "checkToken+当前口令校验(主体XSS-004/CMD-004同机制已验)", "SEQ-C02 反例"),
    ("DVWA", "api(主体AUTH-001)", "R3 gordonb", "dvwa.users",
     "读取/更新用户(low~high)", "已登录", "越权成功", "违反",
     "主体已验:普通用户low/medium/high可改管理数据,impossible拒绝", "主体 wp4-20261019-auth001-*"),
]

# (ID, 不变量, 对象/流程, 检查结论, 证据)
INVARIANTS = [
    ("INV-01", "口令变更必须验证当前口令或等效服务端凭证", "3.1/3.3/3.4改密流",
     "Less-24部分成立(可注入绕过);captcha low/medium违反;csrf low违反", "OPT4-LG-01/02, 主体SQLI-005"),
    ("INV-02", "多步流程的步骤迁移只能由服务端状态判定", "captcha两步流",
     "low违反(不校验);medium违反(校验客户端参数);impossible成立", "OPT4-LG-01a/01b"),
    ("INV-03", "状态变更动作必须绑定会话主体且不可跨站代替发起", "csrf;api",
     "csrf low违反;api侧impossible成立(主体已验)", "OPT4-LG-02; 主体AUTH-001"),
    ("INV-04", "落库数据读回渲染必须按输出上下文编码", "logged-in;xss_s",
     "logged-in违反;xss_s medium的name违反/message成立;impossible成立", "OPT4-SO-01/02"),
    ("INV-05", "落库数据再进入SQL必须参数化(存储时转义不豁免)", "pass_change",
     "违反(主体SQLI-005已验UPDATE汇点);注册写入端成立(反例)", "主体SQLI-005"),
    ("INV-06", "重复提交不产生放大危害(幂等/防重放)", "captcha step=2;csrf",
     "重复覆盖口令类,无放大效应,记录为低危;无防重放机制", "verification_log.csv"),
    ("INV-07", "并发竞态不破坏不变量", "全部",
     "作用域内无可竞争共享计数/余额对象,构造窗口不足→待验证", "optional4/report.md"),
]

# (用例ID, 序列, 回溯样本, 状态)
CASES = [
    ("OPT4-SO-01", "SEQ-A01", "SQLI-005", "待执行"),
    ("OPT4-SO-02", "SEQ-A02", "XSS-005", "待执行"),
    ("OPT4-LG-01a", "SEQ-A03", "新扫描面captcha(low)", "待执行"),
    ("OPT4-LG-01b", "SEQ-A04", "新扫描面captcha(medium)", "待执行"),
    ("OPT4-LG-02", "SEQ-A05", "新扫描面csrf(low)", "待执行"),
    ("SEQ-C01/C02", "反例序列", "XSS-004/CMD-004机制复用", "待执行"),
]


def style_header(ws, widths):
    for cell in ws[1]:
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = WRAP
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "状态矩阵"
    ws.append(["靶场", "模块/代码位置", "角色", "对象", "动作", "前置状态", "后置状态",
               "服务端强制", "判定依据", "关联用例"])
    for row in MATRIX:
        ws.append(list(row))
    style_header(ws, [10, 30, 16, 18, 22, 14, 16, 12, 40, 26])
    for row in ws.iter_rows(min_row=2):
        verdict = row[7].value or ""
        if "违反" in verdict:
            row[7].fill = VIOLATED_FILL
        elif "成立" in verdict:
            row[7].fill = ENFORCED_FILL

    ws2 = wb.create_sheet("业务不变量")
    ws2.append(["ID", "不变量表述", "对象/流程", "检查结论", "证据/用例"])
    for row in INVARIANTS:
        ws2.append(list(row))
    style_header(ws2, [8, 44, 18, 52, 26])

    ws3 = wb.create_sheet("用例映射")
    ws3.append(["用例ID", "请求序列", "回溯样本/扫描面", "执行状态"])
    for row in CASES:
        ws3.append(list(row))
    style_header(ws3, [14, 12, 30, 12])

    wb.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
